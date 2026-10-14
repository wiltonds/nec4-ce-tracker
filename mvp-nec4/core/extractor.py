"""Document ingestion: read text, redact personal data, extract structured fields.

Uses Claude when ANTHROPIC_API_KEY is set; otherwise falls back to simple
rules (regex + keywords) so the demo always runs offline.
"""

import io
import json
import os
import re
from datetime import datetime

EMAIL_RE = re.compile(r"[\w.+-]+@[\w-]+\.[\w.-]+")
PHONE_RE = re.compile(r"(\+?\d[\d\s()-]{8,}\d)")
REF_RE = re.compile(r"\b(EW-\d+|GC-\d+|DOC-\d+|PMI-\d+|CE-\d+|NCE-\d+)\b", re.I)
CH_RE = re.compile(r"\bCH\s?\d{3,5}\b", re.I)
DATE_RES = [
    (re.compile(r"\b(\d{1,2})/(\d{1,2})/(20\d{2})\b"), "%d/%m/%Y"),
    (re.compile(r"\b(20\d{2})-(\d{2})-(\d{2})\b"), "%Y-%m-%d"),
    (re.compile(r"\b(\d{1,2} (?:January|February|March|April|May|June|July|August|September|October|November|December) 20\d{2})\b"), "%d %B %Y"),
]

KEYWORDS = {
    "60.1(5)": ["utility", "utilities", "statutory undertaker", "fibre", "fiber", "diversion", "openreach", "others did not", "not within the times"],
    "60.1(2)": ["access", "possession", "land not available"],
    "60.1(3)": ["not provided", "information not received", "client to provide"],
    "60.1(12)": ["physical condition", "unforeseen", "ground conditions", "collapse", "structural condition", "obstruction"],
    "60.1(1)": ["change to the scope", "not fit for purpose", "scope change", "additional work", "instruction"],
    "60.1(6)": ["no reply", "not replied", "awaiting response", "period for reply"],
    "15.1": ["early warning", "risk", "could delay", "may delay"],
}

MODEL = os.getenv("ANTHROPIC_MODEL", "claude-sonnet-5-5")


def read_text(name: str, data: bytes) -> str:
    """Extract plain text from PDF, DOCX, EML or TXT bytes."""
    lower = name.lower()
    if lower.endswith(".pdf"):
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        return "\n".join((p.extract_text() or "") for p in reader.pages)
    if lower.endswith(".docx"):
        import docx
        d = docx.Document(io.BytesIO(data))
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        return "\n".join(parts)
    if lower.endswith(".eml"):
        from email import policy
        from email.parser import BytesParser
        msg = BytesParser(policy=policy.default).parsebytes(data)
        body = msg.get_body(preferencelist=("plain", "html"))
        return f"Subject: {msg['subject']}\nDate: {msg['date']}\n\n{body.get_content() if body else ''}"
    return data.decode("utf-8", errors="ignore")


def redact(text: str) -> str:
    """Remove e-mails and phone numbers before any processing (data minimisation)."""
    text = EMAIL_RE.sub("[EMAIL REDACTED]", text)
    return PHONE_RE.sub("[PHONE REDACTED]", text)


def _first_date(text: str):
    for rx, fmt in DATE_RES:
        m = rx.search(text)
        if m:
            try:
                return datetime.strptime(m.group(0), fmt).date().isoformat()
            except ValueError:
                continue
    return None


def heuristic_extract(text: str, filename: str) -> dict:
    low = text.lower()
    clauses = [c for c, kws in KEYWORDS.items() if any(k in low for k in kws)]
    refs = REF_RE.findall(text)
    ch = CH_RE.findall(text)
    lines = [l.strip() for l in text.splitlines() if l.strip()]
    subj = next((l.split(":", 1)[1].strip() for l in lines if l.lower().startswith("subject:")), None)
    body = [l for l in lines if not re.match(r"^(from|to|date|ref|subject|cc):", l, re.I)]
    first_line = (subj + " - " if subj else "") + (body[0] if body else filename)
    return {
        "doc_date": _first_date(text),
        "ref": refs[0].upper() if refs else filename,
        "doc_type": "Early Warning" if "early warning" in low else "Correspondence",
        "sender": next((l.split(":", 1)[1].strip() for l in lines if l.lower().startswith("from:")), "Unknown"),
        "summary": first_line[:280],
        "location": ch[0].upper() if ch else "",
        "affected_activity": "",
        "clauses": "; ".join(clauses),
        "potential_ce": any(c.startswith("60.1") for c in clauses),
        "reasoning": "Keyword rules (offline mode). Set ANTHROPIC_API_KEY for AI extraction.",
        "engine": "rules",
    }


PROMPT = """You analyse construction contract correspondence under NEC4 ECC (June 2017).
Return ONLY a JSON object with these keys:
doc_date (YYYY-MM-DD or null), ref (document reference such as EW-95, GC-1385),
doc_type (Early Warning | PM Reply | PM Instruction | Correspondence | Technical Report | CE Notification | Other),
sender (role only, e.g. "Client (Project Manager)", "Utility company", "Contractor" - NEVER a person's name),
summary (max 2 sentences, no personal names - use roles),
location (chainage / km / structure, e.g. "CH 1600"),
affected_activity (the planned activity affected, short),
clauses (semicolon list of NEC4 clauses engaged, e.g. "60.1(5); 61.3"),
potential_ce (true/false: could this support a compensation event for the Contractor?),
reasoning (1-2 sentences why),
package_id (one id from the package list below),
topic_id (the existing topic this fact belongs to, from the topic list below, or "NEW"),
work_package (the work package / structure / section affected, short).

Packages (id: name):
{packages}

Existing topics (id: title):
{topics}

Programme activities for context:
{activities}

Document:
\"\"\"{text}\"\"\""""


def ai_extract(text: str, filename: str, activities: str, packages: str = "", topics: str = "") -> dict:
    import anthropic
    client = anthropic.Anthropic()
    msg = client.messages.create(
        model=MODEL,
        max_tokens=800,
        messages=[{"role": "user", "content": PROMPT.format(activities=activities[:6000], packages=packages,
                                                                  topics=topics[:6000], text=text[:30000])}],
    )
    raw = msg.content[0].text
    raw = raw[raw.find("{"): raw.rfind("}") + 1]
    out = json.loads(raw)
    out["engine"] = MODEL
    return out


def extract(name: str, data: bytes, activities: str = "", packages: str = "", topics: str = "") -> tuple[dict, str]:
    text = redact(read_text(name, data))
    if os.getenv("ANTHROPIC_API_KEY"):
        try:
            return ai_extract(text, name, activities, packages, topics), text
        except Exception as e:  # keep the demo running
            res = heuristic_extract(text, name)
            res["reasoning"] = f"AI failed ({e.__class__.__name__}); used keyword rules."
            return res, text
    return heuristic_extract(text, name), text


def refine_narrative(draft: str) -> str:
    """Optional: ask Claude to tighten a programme narrative without changing facts or figures."""
    import anthropic
    client = anthropic.Anthropic()
    msg = client.messages.create(
        model=MODEL,
        max_tokens=3000,
        messages=[{"role": "user", "content":
                   "You are an NEC4 planner writing the delay narrative for a Clause 32 revised programme. "
                   "Improve clarity and flow of this narrative. Keep every date, reference, activity code and "
                   "number exactly as given, do not invent facts, use roles not personal names, keep the section "
                   "structure. Return only the narrative.\n\n" + draft}],
    )
    return msg.content[0].text


def refine_notice(draft: str) -> str:
    """Optional: ask Claude to tighten a draft CE notice."""
    import anthropic
    client = anthropic.Anthropic()
    msg = client.messages.create(
        model=MODEL,
        max_tokens=2000,
        messages=[{"role": "user", "content":
                   "You are an NEC4 commercial manager. Improve this Contractor's compensation event "
                   "notification: keep facts and dates exactly as given, keep it factual and firm, "
                   "do not invent facts, use roles not personal names. Return only the notice.\n\n" + draft}],
    )
    return msg.content[0].text
